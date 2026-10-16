"""
Lecture d'un import d'élèves (WT06, C19) : fichier CSV ou Excel (.xlsx), 1 Mo et 200 lignes au plus, ou lignes déjà
découpées par le front (liste collée, lignes corrigées renvoyées). Colonnes : Prénom, Nom, Téléphone du parent, Lien
de parenté ; une ligne d'en-tête est reconnue et ignorée. Le serveur seul juge les doublons, les numéros, les quotas
et l'appartenance à la classe : le front ne valide que la forme.
"""

import csv
import io
import zipfile

from libs.errors import ApiError
from libs.phone import normalize_phone

from .learner_invitations import ImportResult, Invitee, fold

MAX_ROWS = 200
MAX_BYTES = 1024 * 1024
#: Colonnes lues dans un classeur : les 4 attendues et une marge pour signaler une ligne mal formée.
MAX_COLUMNS = 8
#: Lignes parcourues au plus dans un classeur, vides comprises (un classeur compressé peut en déclarer des millions).
MAX_SCANNED = 5000
MAX_NAME = 150
COLUMNS = ("first_name", "last_name", "parent_phone", "relationship")
HEADER_WORDS = {"prenom", "firstname", "first name"}


class _Semicolon(csv.excel):
    """Séparateur des tableurs français, quand le fichier ne permet pas de le deviner."""

    delimiter = ";"


def _cell(value) -> str:
    """Un numéro saisi dans Excel arrive parfois en nombre décimal (97000001.0)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value if value is not None else "").strip()


def _read_xlsx(content: bytes) -> list[list]:
    """
    Lecture en flux, bornée (bombe de décompression) : `MAX_COLUMNS` colonnes, arrêt dès `MAX_ROWS + 2` lignes non
    vides (en-tête compris) ou `MAX_SCANNED` lignes parcourues. Le XML est lu par `defusedxml` (openpyxl l'utilise
    dès qu'il est installé).
    """
    from openpyxl import load_workbook

    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        rows, filled = [], 0
        for scanned, row in enumerate(
            sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True), start=1
        ):
            rows.append(list(row))
            filled += any(_cell(cell) for cell in row)
            if filled >= MAX_ROWS + 2 or scanned >= MAX_SCANNED:
                raise ApiError("IMPORT_TOO_LARGE")
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise ApiError("IMPORT_INVALID_FILE") from error
    return rows


def _read_csv(content: bytes) -> list[list]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = _Semicolon
    if "\x00" in text:  # accepté par `csv` depuis Python 3.11, refusé par PostgreSQL
        raise ApiError("IMPORT_INVALID_FILE")
    try:
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error as error:  # champ démesuré, guillemets incohérents…
        raise ApiError("IMPORT_INVALID_FILE") from error


def read_file(upload) -> list[tuple[int, list]]:
    """Lignes non vides du fichier, numérotées comme dans le tableur (1 = première ligne)."""
    if upload.size > MAX_BYTES:
        raise ApiError(
            "IMPORT_TOO_LARGE", "Le fichier dépasse 1 Mo : découpez-le en plusieurs imports."
        )
    content = upload.read()
    rows = _read_xlsx(content) if content[:2] == b"PK" else _read_csv(content)
    lines = [
        (number, row)
        for number, row in enumerate(rows, start=1)
        if any(_cell(cell) for cell in row)
    ]
    if lines and fold(lines[0][1][0]) in HEADER_WORDS:
        lines = lines[1:]
    return lines


def check_size(lines) -> None:
    if len(lines) > MAX_ROWS:
        raise ApiError("IMPORT_TOO_LARGE")


def parse(lines, *, relationships, dial_code: str, phone_valid) -> tuple[list, ImportResult]:
    """
    Valide la forme de chaque ligne. Renvoie les lignes prêtes (`(ligne, Invitee, données)`) et le résultat où sont
    déjà consignées les lignes mal formées, aux numéros invalides et en double.
    """
    by_name = {}
    for relationship in relationships:
        by_name[fold(relationship.code)] = relationship
        by_name[fold(relationship.label)] = relationship
    result, ready, seen = ImportResult(), [], set()
    for line, cells in lines:
        cells = [_cell(cell) for cell in cells]
        # Ligne courte : les colonnes manquantes sont absentes des données renvoyées
        values = data = dict(zip(COLUMNS, cells[:4], strict=False))
        relationship = by_name.get(fold(values.get("relationship")))
        if (
            len(cells) < 4
            or not all(cells[:4])
            or any(c for c in cells[4:])
            or relationship is None
            or len(values["first_name"]) > MAX_NAME
            or len(values["last_name"]) > MAX_NAME
        ):
            result.fail(line, "MALFORMED_ROW", data)
            continue
        phone = normalize_phone(values["parent_phone"], dial_code)
        if phone is None or not phone_valid(phone):
            result.fail(line, "PHONE_INVALID", data)
            continue
        invitee = Invitee(values["first_name"], values["last_name"], phone, relationship)
        if invitee.key() in seen:
            result.fail(line, "DUPLICATE_ROW", data)
            continue
        seen.add(invitee.key())
        ready.append((line, invitee, data))
    return ready, result
